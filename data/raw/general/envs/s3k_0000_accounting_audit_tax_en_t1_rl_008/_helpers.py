"""共享验证助手 — 出题时的 gt_code 执行与交付后的 verify.py 共用同一套原语。

由 make_tasks.py 写入 <wd>/tasks/_helpers.py;verify.py 从父目录 import。
"""
from __future__ import annotations

import hashlib
import json  # noqa: F401  (gt/check 代码常用)
import re
import subprocess
import tempfile
from pathlib import Path


def extract_text(path) -> str:
    """docx/pptx/xlsx → 纯文本(段落/表格逐行;xlsx 用重算后的计算值)。"""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [par.text for par in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(p))
        parts = []
        for i, s in enumerate(prs.slides, 1):
            parts.append(f"# slide {i}")
            for sh in s.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False):
                    for row in sh.table.rows:
                        parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(recalc_xlsx(p)), data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"# {ws.title}")
            for r in ws.iter_rows(values_only=True):
                if any(x is not None for x in r):
                    parts.append(" | ".join("" if x is None else str(x) for x in r))
        wb.close()
        return "\n".join(parts)
    if ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(str(p))
            return "\n".join(doc[i].get_text() for i in range(doc.page_count))
        except Exception:
            try:
                import pdfplumber
                with pdfplumber.open(str(p)) as pdf:
                    return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
            except Exception:
                return ""
    if ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"):
        return ""   # 图片无文本层(视觉判分走 judge_quality vision;需文字用 OCR,当前不做)
    return p.read_text(errors="ignore")   # csv/tsv/txt/md 等纯文本直接读


def recalc_xlsx(path) -> Path:
    """soffice 重算出带公式缓存值的副本(openpyxl 产物无缓存值);失败返回原路径。"""
    src = Path(path)
    key = hashlib.sha1(str(src.resolve()).encode()).hexdigest()[:10]
    out_dir = Path(tempfile.gettempdir()) / f"verify_rc_{key}"
    out = out_dir / src.name
    if out.exists() and out.stat().st_mtime >= src.stat().st_mtime:
        return out
    out_dir.mkdir(exist_ok=True)
    try:
        subprocess.run(["soffice", "--headless",
                        f"-env:UserInstallation=file:///tmp/soffice_v_{key}",
                        "--convert-to", "xlsx", "--outdir", str(out_dir), str(src)],
                       capture_output=True, timeout=240)
    except Exception:
        return src
    return out if out.exists() else src


def load_sheet(path, sheet, recalc=True):
    """打开某 sheet(默认 soffice 重算后 data_only,公式读到计算值)。"""
    import openpyxl
    p = recalc_xlsx(path) if recalc else Path(path)
    wb = openpyxl.load_workbook(str(p), data_only=True)
    return wb[sheet]


def _ev_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:g}"
    return str(v)


def _ev_assemble(blocks, budget: int):
    """blocks=[(kind,text)] → (text, truncated)。超预算时表格块优先整块保留,
    被裁的块合并成一条显式【已截断】标记(禁止静默丢内容)。"""
    bs = []
    for kind, txt in blocks:
        if len(txt) > budget:   # 单块就超预算:块内留头部+标记
            head = max(0, budget - 80)
            txt = txt[:head] + f"\n【已截断:本块原 {len(txt)} 字符,仅保留前 {head}】"
        bs.append((kind, txt))
    if sum(len(t) + 1 for _, t in bs) <= budget:
        return "\n".join(t for _, t in bs), False
    keep, used = set(), 0
    for i, (k, t) in enumerate(bs):
        if k == "table" and used + len(t) + 1 <= budget:
            keep.add(i); used += len(t) + 1
    for i, (k, t) in enumerate(bs):
        if i not in keep and used + len(t) + 1 <= budget:
            keep.add(i); used += len(t) + 1
    parts, cut_n, cut_ch = [], 0, 0
    for i, (k, t) in enumerate(bs):
        if i in keep:
            if cut_n:
                parts.append(f"【已截断:省略 {cut_n} 块共 {cut_ch} 字符(表格优先保留)】")
                cut_n = cut_ch = 0
            parts.append(t)
        else:
            cut_n += 1; cut_ch += len(t)
    if cut_n:
        parts.append(f"【已截断:省略 {cut_n} 块共 {cut_ch} 字符(表格优先保留)】")
    return "\n".join(parts), True


def _ev_zip_count(path, prefix: str) -> int:
    import zipfile
    try:
        with zipfile.ZipFile(str(path)) as z:
            return sum(1 for n in z.namelist() if n.startswith(prefix))
    except Exception:
        return 0


def extract_evidence(path, budget: int = 20000) -> dict:
    """交付物 → LLM judge 证据文本(可审计):保文档序、表格网格带行号/列标+来源标注、
    未抽取内容清单、结构优先的显式截断(任何裁剪必留【已截断】标记,judge 据此区分
    "内容不存在"与"没给我看")。返回 {"text","loss","truncated","raw_chars"}。
    与 extract_text 分工:extract_text 供 rule 代码机器解析(输出格式保持稳定不动);
    本函数只服务 llm 判分证据,可持续增强不影响 rule 兼容。"""
    p = Path(path)
    ext = p.suffix.lower()
    blocks: list = []          # (kind: "table"|"text", text)
    loss: list[str] = []

    if ext == ".docx":
        import docx
        from docx.table import Table as _DxTable
        from docx.text.paragraph import Paragraph as _DxPara
        d = docx.Document(str(p))
        ti = 0
        for child in d.element.body.iterchildren():   # 按 body 原始顺序:段落/表格交错
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                par = _DxPara(child, d)
                txt = par.text.strip()
                if not txt:
                    continue
                style = getattr(par.style, "name", "") or ""
                if style.startswith("Heading"):
                    txt = f"[H{style.split()[-1]}] {txt}"
                blocks.append(("text", txt))
            elif tag == "tbl":
                ti += 1
                tbl = _DxTable(child, d)
                lines = [f"[docx:表{ti} {len(tbl.rows)}行x{len(tbl.columns)}列]"]
                for ri, row in enumerate(tbl.rows, 1):
                    lines.append(f"r{ri}: " + " | ".join(c.text.strip() for c in row.cells))
                blocks.append(("table", "\n".join(lines)))
        try:   # 文本框(python-docx 不在 body 流里,单独抽)
            tb = [t.text for t in d.element.body.xpath(".//w:txbxContent//w:t")
                  if (t.text or "").strip()]
            if tb:
                blocks.append(("text", "[文本框] " + " ".join(tb)))
        except Exception:
            pass
        if len(d.inline_shapes):
            loss.append(f"{len(d.inline_shapes)} 张内嵌图片")
        nch = _ev_zip_count(p, "word/charts/chart")
        if nch:
            loss.append(f"{nch} 个图表")

    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(p))
        n_pic = n_chart = 0

        def _walk(shapes, si):
            nonlocal n_pic, n_chart
            for sh in shapes:
                if getattr(sh, "shape_type", None) is not None and hasattr(sh, "shapes"):
                    _walk(sh.shapes, si)   # 组合形状递归下钻
                    continue
                if getattr(sh, "has_chart", False):
                    n_chart += 1
                try:
                    if int(getattr(sh, "shape_type", 0) or 0) == 13:   # PICTURE
                        n_pic += 1
                except Exception:
                    pass
                if getattr(sh, "has_text_frame", False):
                    txt = sh.text_frame.text.strip()
                    if txt:
                        blocks.append(("text", txt))
                if getattr(sh, "has_table", False):
                    lines = [f"[pptx:slide{si} 表格 {len(sh.table.rows)}行x{len(sh.table.columns)}列]"]
                    for ri, row in enumerate(sh.table.rows, 1):
                        lines.append(f"r{ri}: " + " | ".join(c.text.strip() for c in row.cells))
                    blocks.append(("table", "\n".join(lines)))

        for i, s in enumerate(prs.slides, 1):
            blocks.append(("text", f"[pptx:slide {i}]"))
            _walk(s.shapes, i)
            try:
                if s.has_notes_slide:
                    nt = s.notes_slide.notes_text_frame.text.strip()
                    if nt:
                        blocks.append(("text", f"[slide {i} 备注] {nt}"))
            except Exception:
                pass
        if n_pic:
            loss.append(f"{n_pic} 张图片")
        if n_chart:
            loss.append(f"{n_chart} 个图表")

    elif ext == ".xlsx":
        import openpyxl
        from openpyxl.utils import get_column_letter
        rp = recalc_xlsx(p)
        try:
            wb = openpyxl.load_workbook(str(rp), data_only=True)
            ro = False
        except Exception:
            wb = openpyxl.load_workbook(str(rp), data_only=True, read_only=True)
            ro = True
        ROWS_CAP = 1500
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            ncol = 0
            for r in rows:
                for j in range(len(r) - 1, -1, -1):
                    if r[j] is not None:
                        ncol = max(ncol, j + 1)
                        break
            ncol = min(ncol, 40)
            lines = [f"[xlsx:sheet「{ws.title}」{len(rows)}行x{ncol}列]",
                     "   " + " | ".join(get_column_letter(j + 1) for j in range(ncol))]
            shown = 0
            for ri, r in enumerate(rows, 1):
                if not any(x is not None for x in r[:ncol]):
                    continue   # 空行跳过但行号保留在 r{n} 里,不错位
                lines.append(f"r{ri}: " + " | ".join(_ev_cell(x) for x in r[:ncol]))
                shown += 1
                if shown >= ROWS_CAP:
                    lines.append(f"【已截断:sheet「{ws.title}」剩余 {len(rows) - ri} 行未列出】")
                    break
            if not ro:
                try:
                    mr = [str(x) for x in ws.merged_cells.ranges][:20]
                    if mr:
                        lines.append("[合并单元格] " + ", ".join(mr))
                except Exception:
                    pass
            blocks.append(("table", "\n".join(lines)))
        wb.close()
        nch = _ev_zip_count(rp, "xl/charts/chart")
        if nch:
            loss.append(f"{nch} 个图表")
        loss.append("单元格显示格式未保留(值为重算后计算值)")

    elif ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(str(p))
            n_img = 0
            for i in range(doc.page_count):
                pg = doc[i]
                txt = pg.get_text().strip()
                blocks.append(("text", f"[pdf:page {i + 1}]\n{txt}"))
                try:
                    n_img += len(pg.get_images())
                except Exception:
                    pass
            if n_img:
                loss.append(f"{n_img} 张图片")
        except Exception as e:
            loss.append(f"pdf 解析失败: {e}")

    elif ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"):
        loss.append("图片文件无文本层(该判定需 judge:vision)")

    else:   # csv/tsv/txt/md 等纯文本
        blocks.append(("text", p.read_text(errors="ignore")))

    raw = sum(len(t) for _, t in blocks)
    text, truncated = _ev_assemble(blocks, budget)
    return {"text": text, "loss": loss, "truncated": truncated, "raw_chars": raw}


def _num(x):
    if isinstance(x, bool):
        return None
    try:
        return float(str(x).replace(",", "").replace("%", ""))
    except (TypeError, ValueError):
        return None


def norm_label(s) -> str:
    """标签归一:去空白/换行、全角→半角、去括号内容差异容忍(仅去空白与全角),小写。"""
    if s is None:
        return ""
    t = str(s)
    t = t.translate(str.maketrans("０１２３４５６７８９（）：％．－", "0123456789():%.-"))
    return re.sub(r"\s+", "", t).lower()


def find_col(header_cells, *keywords, required=True):
    """模糊列定位:返回首个「归一后包含全部 keyword」的列 index;找不到 required 则 raise。
    用于读 agent 产出的交付物(列名可能有措辞差异),避免硬编码列名/索引崩溃。"""
    kws = [norm_label(k) for k in keywords]
    for i, c in enumerate(header_cells):
        lab = norm_label(c)
        if lab and all(k in lab for k in kws):
            return i
    if required:
        raise KeyError(f"find_col: 没有列同时含 {keywords};表头={[str(c)[:20] for c in header_cells]}")
    return None


def find_header_row(rows, *keywords, max_scan=15):
    """在前 max_scan 行里找表头行(归一后包含全部 keyword 中每个词的某列)。
    返回 (row_idx, header_cells);找不到 raise。rows=list(iter_rows(values_only=True))。"""
    kws = [norm_label(k) for k in keywords]
    for i, row in enumerate(rows[:max_scan]):
        labs = [norm_label(c) for c in row]
        if all(any(k in lab for lab in labs if lab) for k in kws):
            return i, list(row)
    raise KeyError(f"find_header_row: 前 {max_scan} 行找不到含 {keywords} 的表头行")


def find_row(rows, label, col=0):
    """精确等值行定位(归一比较,防「广西」子串误伤「广西较基准差异」)。找不到 raise。"""
    want = norm_label(label)
    for row in rows:
        if col < len(row) and norm_label(row[col]) == want:
            return row
    raise KeyError(f"find_row: 第{col}列无值精确等于 {label!r}")


_ERR_CELLS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NULL!", "#NUM!")


def file_healthy(path) -> tuple[bool, str]:
    """交付物结构完整性(供 check.py 作基础层子项):文件能打开、xlsx 无错误单元格。
    返回 (ok, reason)。不 raise——坏文件是可判事实,不是异常。"""
    p = Path(path)
    if not p.exists():
        return False, "文件不存在"
    ext = p.suffix.lower()
    try:
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(recalc_xlsx(p)), data_only=True, read_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        if isinstance(v, str) and v.strip() in _ERR_CELLS:
                            wb.close()
                            return False, f"含错误单元格 {v.strip()}(sheet {ws.title})"
            wb.close()
        elif ext == ".docx":
            import docx
            docx.Document(str(p))
        elif ext == ".pptx":
            from pptx import Presentation
            Presentation(str(p))
        else:
            p.read_bytes()
    except Exception as e:
        return False, f"打开失败/损坏: {e}"
    return True, "ok"


def locate_deliverable(ws, *name_keywords, exts=(".xlsx", ".docx", ".pptx", ".csv")):
    """鲁棒定位 agent 产出的交付物:优先文件名归一含全部 keyword,退而取该类型唯一新文件。
    防"内容对但文件名没精确匹配"误判 0。找不到返回 None(由 check 记基础层失分,不崩)。"""
    ws = Path(ws)
    cands = [p for p in ws.rglob("*") if p.is_file() and p.suffix.lower() in exts
             and not any(part.startswith(".") for part in p.relative_to(ws).parts)]
    if name_keywords:
        kws = [norm_label(k) for k in name_keywords]
        hit = [p for p in cands if all(k in norm_label(p.stem) for k in kws)]
        if hit:
            return sorted(hit, key=lambda p: len(p.stem))[0]
    return cands[0] if len(cands) == 1 else None


def _canon(s) -> str:
    n = _num(s)
    if n is not None:
        return f"{round(n, 2):g}"
    return re.sub(r"\s+", "", str(s)).strip().lower()


def deep_score(exp, got):
    """叶级部分分,题间分值可比的关键原语:
    - 数值:浮点容差 max(0.01, 0.5%);其余标量归一后精确比对
    - 标量列表:多重集 Jaccard 部分分(对一半 ≈0.5,不再一票制)
    - 对象列表:按**首字段**(应为唯一标识,出题规范要求)对齐匹配后逐项递归
      (cohort 对 4/5 只错 1 项得 ~0.8,不因错位/排序全灭)
    - dict 按键递归;score = 全部检查点的平均
    """
    checks = []   # (score01, desc)

    def walk(e, g, path=""):
        if isinstance(e, dict):
            if not isinstance(g, dict):
                checks.append((0.0, f"{path}: 类型应为 object"))
                return
            for k, v in e.items():
                walk(v, g.get(k), f"{path}.{k}")
        elif isinstance(e, list):
            if not isinstance(g, list):
                checks.append((0.0, f"{path}: 类型应为 list"))
                return
            if all(not isinstance(x, (dict, list)) for x in e) and \
               all(not isinstance(x, (dict, list)) for x in g):
                from collections import Counter
                ce, cg = Counter(_canon(x) for x in e), Counter(_canon(x) for x in g)
                inter = sum((ce & cg).values())
                union = sum((ce | cg).values())
                s = (inter / union) if union else 1.0
                checks.append((s, f"{path}: 集合重合 {inter}/{union} "
                                  f"exp={sorted(ce)[:5]} got={sorted(cg)[:5]}"))
            elif e and all(isinstance(x, dict) for x in e) and \
                    g and all(isinstance(x, dict) for x in g):
                key = next(iter(e[0].keys()), None)
                ekeys = [_canon(ei.get(key)) for ei in e]
                if len(set(ekeys)) != len(ekeys):   # 出题 bug:首字段须唯一标识,重复则对齐不可靠
                    checks.append((0.0, f"{path}: expected 首字段 {key!r} 有重复值,无法可靠对齐"
                                        f"(出题需保证首字段唯一)"))
                gmap: dict = {}
                for x in g:
                    gmap.setdefault(_canon(x.get(key)), x)
                if len(e) != len(g):
                    checks.append((0.0, f"{path}: 长度 {len(e)}≠{len(g)}"))
                for i, ei in enumerate(e):
                    gi = gmap.get(_canon(ei.get(key)))
                    if gi is None:
                        checks.append((0.0, f"{path}[{i}]({key}={ei.get(key)!r}): 该项缺失"))
                    else:
                        walk(ei, gi, f"{path}[{i}]")
                # 超集罚:got 里 expected 之外的 key 逐个记 0(否则照抄完整池只损 1 条长度检查,
                # 20+1 里得 0.95;逐个记 0 后懒答案 near-0,near-miss 仍有部分分——两头都对)
                extra = set(gmap) - set(ekeys)
                for k in sorted(extra)[:200]:
                    checks.append((0.0, f"{path}: 多余项 {key}={k!r}(不在目标集内)"))
            else:
                if len(e) != len(g):
                    checks.append((0.0, f"{path}: 长度 {len(e)}≠{len(g)}"))
                for i in range(min(len(e), len(g))):
                    walk(e[i], g[i], f"{path}[{i}]")
        else:
            ne, ng = _num(e), _num(g)
            if ne is not None and ng is not None:
                ok = abs(ne - ng) <= max(0.01, abs(ne) * 0.005)
            else:
                ok = _canon(e) == _canon(g)
            checks.append((1.0 if ok else 0.0, f"{path}: exp={e!r} got={g!r}"))

    walk(exp, got)
    if not checks:
        return 0.0, ["expected 为空"]
    fails = [d for s, d in checks if s < 0.999]
    return sum(s for s, _ in checks) / len(checks), fails[:10]


# ---------------------------------------------------------------- 源守恒 gate(写足迹反推,自动生成)

def _sha_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def ws_fingerprints(ws) -> dict:
    """workspace 全部文件的内容指纹 {相对路径: sha256}(隐藏目录跳过)。
    源守恒 gate 用:出题时按 solve 写足迹反推保护集,判分时比对。"""
    ws = Path(ws)
    out = {}
    for p in sorted(ws.rglob("*")):
        if not p.is_file():
            continue
        rel = p.relative_to(ws).as_posix()
        if any(part.startswith(".") for part in rel.split("/")):
            continue
        try:
            out[rel] = _sha_bytes(p.read_bytes())
        except Exception:
            pass
    return out


def db_fingerprints(env_root) -> dict:
    """system/<mcp>/state.db 每表**逻辑指纹** {mcp: {table: sha256}}:行内容排序后 hash
    (行序无关、不受 WAL/文件级差异影响);跳过 sqlite_% 内部表。env_root=含 system/ 的根。"""
    import sqlite3
    out: dict = {}
    sysdir = Path(env_root) / "system"
    if not sysdir.is_dir():
        return out
    for dbf in sorted(sysdir.glob("*/state.db")):
        try:
            con = sqlite3.connect(f"file:{dbf}?mode=ro", uri=True)
            tabs = [r[0] for r in con.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")]
            th = {}
            for t in tabs:
                rows = sorted(_sha_bytes(repr(r).encode()) for r in con.execute(f'SELECT * FROM "{t}"'))
                th[t] = _sha_bytes("|".join(rows).encode())
            out[dbf.parent.name] = th
            con.close()
        except Exception:
            pass
    return out


def protect_diff(ws, protect) -> list:
    """源守恒 gate 比对。protect={'files':{rel:sha},'db_tables':{mcp:{table:sha}}} =
    出题时实测的「solve 不触碰」集合(合法解法的定义足迹,天然零误杀)。
    返回违规清单(空=通过):保护文件被改/删、保护表内容变化;agent **新增**文件不罚。"""
    bad = []
    ws = Path(ws)
    cur = ws_fingerprints(ws)
    for rel, h in (protect.get("files") or {}).items():
        c = cur.get(rel)
        if c is None:
            bad.append(f"保护文件被删除: {rel}")
        elif c != h:
            bad.append(f"保护文件被篡改: {rel}")
    ptabs = protect.get("db_tables") or {}
    if ptabs:
        dbs = db_fingerprints(Path(ws).resolve().parent)
        for mcp, tabs in ptabs.items():
            curt = dbs.get(mcp) or {}
            for t, h in tabs.items():
                c = curt.get(t)
                if c is None:
                    bad.append(f"保护表缺失: {mcp}.{t}")
                elif c != h:
                    bad.append(f"保护表被改动: {mcp}.{t}")
    return bad


# ---------------------------------------------------------------- DB 原语(db_gen 混合环境)

def db_path(ws, mcp) -> Path:
    """业务系统 DB 路径:<env>/system/<mcp>/state.db(ws=workspace 根,system 与其同级)。
    找不到 raise——DB 题的环境必须带 system/(出题/验证副本都会连带拷贝)。"""
    p = Path(ws).resolve().parent / "system" / mcp / "state.db"
    if not p.exists():
        raise FileNotFoundError(f"state.db 不存在: {p}(环境副本没带 system/?)")
    return p


def db_query(ws, mcp, sql, params=()):
    """只读 SQL 查询 → list[dict]。gt/check 访问 DB 的唯一通道(含 mutation 题查后态)。
    gt_code/check_code 禁写库(静态检查拦 db_execute)——verify 必须无副作用。"""
    import sqlite3
    con = sqlite3.connect(f"file:{db_path(ws, mcp)}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in con.execute(sql, tuple(params)).fetchall()]
    finally:
        con.close()


def db_scalar(ws, mcp, sql, params=()):
    """只读 SQL 返回首行首列标量(聚合值核对用);无结果返回 None。"""
    rows = db_query(ws, mcp, sql, params)
    if not rows:
        return None
    return next(iter(rows[0].values()))


def db_execute(ws, mcp, sql, params=()):
    """写库(INSERT/UPDATE/DELETE),返回受影响行数。**仅 solve.py 正解脚本可用**
    (自证 mutation 题可满分);gt/check 用它会被出题门静态拒绝。"""
    import sqlite3
    con = sqlite3.connect(str(db_path(ws, mcp)))
    try:
        cur = con.execute(sql, tuple(params))
        con.commit()
        return cur.rowcount
    finally:
        con.close()
